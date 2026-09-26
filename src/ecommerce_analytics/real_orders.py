"""Order-level analysis for public transaction exports without item/customer IDs.

This profile deliberately does not coerce an order-grain export into the richer
item-level schema used by the demo. Metrics here stay at order, date, payment,
refund, and region grain.
"""

from __future__ import annotations

import html
import io
import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .pipeline import numeric_price, parse_date

ALIASES = {
    "订单编号": "订单号",
    "订单号": "订单号",
    "总金额": "订单金额",
    "订单金额": "订单金额",
    "买家实际支付金额": "实付金额",
    "实付金额": "实付金额",
    "收货地址": "收货省份",
    "收货省份": "收货省份",
    "订单创建时间": "下单时间",
    "下单时间": "下单时间",
    "订单付款时间": "付款时间",
    "付款时间": "付款时间",
    "退款金额": "退款金额",
}
REQUIRED = ["订单号", "订单金额", "实付金额", "下单时间"]
OPTIONAL = ["收货省份", "付款时间", "退款金额"]
DATE_COLUMNS = ["下单时间", "付款时间"]
MONEY_COLUMNS = ["订单金额", "实付金额", "退款金额"]


def _normalize_headers(frame: pd.DataFrame) -> pd.DataFrame:
    renamed = frame.copy()
    renamed.columns = [ALIASES.get(str(column).replace("\ufeff", "").strip(), str(column).strip()) for column in frame.columns]
    duplicated = renamed.columns[renamed.columns.duplicated()].tolist()
    if duplicated:
        raise ValueError(f"订单级数据存在重复字段：{', '.join(map(str, duplicated))}")
    missing = [column for column in REQUIRED if column not in renamed.columns]
    if missing:
        raise ValueError(
            "无法识别为订单级导出；缺少字段：" + "、".join(missing)
            + "。需至少包含订单号、订单金额、实付金额、下单时间。"
        )
    keep = REQUIRED + [column for column in OPTIONAL if column in renamed.columns]
    return renamed.loc[:, keep].copy()


def read_order_level_stream(stream: io.IOBase, suffix: str) -> pd.DataFrame:
    """Read a CSV/XLSX order-grain export and normalize known Chinese headers."""
    suffix = suffix.lower()
    if suffix == ".csv":
        content = stream.read()
        if isinstance(content, str):
            content = content.encode("utf-8")
        last_error = None
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                frame = pd.read_csv(io.BytesIO(content), encoding=encoding, dtype="string")
                break
            except UnicodeDecodeError as error:
                last_error = error
        else:
            raise ValueError("CSV 字符编码无法识别（已尝试 UTF-8 与 GB18030）。") from last_error
    elif suffix in {".xlsx", ".xls"}:
        frame = pd.read_excel(stream, dtype="string")
    else:
        raise ValueError("输入必须是 CSV 或 Excel 文件。")
    return _normalize_headers(frame)


def read_order_level(path: Path) -> pd.DataFrame:
    path = Path(path)
    with path.open("rb") as stream:
        return read_order_level_stream(stream, path.suffix)


def clean_order_level(frame: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, Any]]:
    """Clean supported order-level fields without fabricating absent dimensions."""
    raw = _normalize_headers(frame)
    optional_supplied = [column for column in OPTIONAL if column in raw.columns]
    result = raw.copy()
    for column in result.columns:
        result[column] = result[column].astype("string").str.strip().replace("", pd.NA)
    result["订单号"] = result["订单号"].astype("string")
    for column in MONEY_COLUMNS:
        if column in result.columns:
            result[column] = numeric_price(result[column])
    for column in DATE_COLUMNS:
        if column in result.columns:
            result[column] = result[column].map(parse_date)
    if "收货省份" in result.columns:
        result["收货省份"] = result["收货省份"].astype("string").str.strip()

    duplicate_rows = int(result["订单号"].duplicated().sum())
    result = result.drop_duplicates(subset=["订单号"], keep="first").reset_index(drop=True)
    source_totals = {column: round(float(result[column].sum()), 2) for column in MONEY_COLUMNS if column in result}
    log = {
        "dataset_type": "order_level_public_export",
        "raw_rows": int(len(raw)),
        "unique_order_ids": int(raw["订单号"].nunique()),
        "duplicate_rows_removed": duplicate_rows,
        "rows_after_deduplication": int(len(result)),
        "missing_order_ids": int(result["订单号"].isna().sum()),
        "unparseable_or_missing_order_dates": int(result["下单时间"].isna().sum()),
        "unparseable_or_missing_payment_dates": int(result["付款时间"].isna().sum()) if "付款时间" in result else None,
        "missing_or_unparseable_money": {
            column: int(result[column].isna().sum()) for column in MONEY_COLUMNS if column in result
        },
        "optional_columns_supplied": optional_supplied,
        "optional_columns_absent": [column for column in OPTIONAL if column not in optional_supplied],
        "source_totals_after_deduplication": source_totals,
        "rules": [
            "Trim header and cell whitespace; treat blank strings as missing.",
            "Parse supported timestamps and monetary amounts; preserve unparseable values as missing.",
            "Deduplicate by order ID and keep the first row; no duplicates were removed when IDs are unique.",
            "Do not infer item, customer, traffic, logistics, or explicit order-status fields.",
        ],
    }
    return result, log


def _money_sum(series: pd.Series) -> float:
    return round(float(series.sum()), 2)


def analyze_order_level(cleaned: pd.DataFrame) -> dict[str, Any]:
    """Compute order, payment, refund, temporal, and geographic metrics."""
    total_orders = int(cleaned["订单号"].nunique())
    positive_paid_mask = cleaned["实付金额"].fillna(0).gt(0)
    has_payment_time = cleaned["付款时间"].notna() if "付款时间" in cleaned.columns else pd.Series(False, index=cleaned.index)
    refund_available = "退款金额" in cleaned.columns and cleaned["退款金额"].notna().any()
    refund_values = cleaned["退款金额"].fillna(0) if "退款金额" in cleaned else pd.Series(0.0, index=cleaned.index)
    refund_mask = refund_values.gt(0)
    order_amount = _money_sum(cleaned["订单金额"].dropna())
    paid_amount = _money_sum(cleaned["实付金额"].dropna())
    refund_amount = _money_sum(refund_values)

    dated = cleaned[cleaned["下单时间"].notna()].copy()
    dated["下单日期"] = dated["下单时间"].dt.strftime("%Y-%m-%d")
    dated["_has_payment_time"] = has_payment_time.loc[dated.index].values
    dated["_positive_paid"] = positive_paid_mask.loc[dated.index].values
    dated["_refund_amount"] = refund_values.loc[dated.index].values
    daily = dated.groupby("下单日期", as_index=False).agg(
        orders=("订单号", "nunique"), order_amount=("订单金额", "sum"), paid_amount=("实付金额", "sum"),
        payment_timestamp_orders=("_has_payment_time", "sum"),
        positive_actual_payment_orders=("_positive_paid", "sum"),
        refund_amount=("_refund_amount", "sum"),
    )
    for column in ["order_amount", "paid_amount", "refund_amount"]:
        daily[column] = daily[column].fillna(0).round(2)

    region: list[dict[str, Any]] = []
    if "收货省份" in cleaned.columns and cleaned["收货省份"].notna().any():
        regional = cleaned[cleaned["收货省份"].notna()].groupby("收货省份", as_index=False).agg(
            orders=("订单号", "nunique"), order_amount=("订单金额", "sum"),
            paid_amount=("实付金额", "sum"), refund_amount=("退款金额", "sum") if "退款金额" in cleaned else ("订单号", "size"),
        )
        if "退款金额" not in cleaned:
            regional["refund_amount"] = 0.0
        regional = regional.sort_values(["paid_amount", "orders"], ascending=False)
        region = [
            {"收货省份": str(row["收货省份"]), "orders": int(row["orders"]),
             "order_amount": round(float(row["order_amount"]), 2), "paid_amount": round(float(row["paid_amount"]), 2),
             "refund_amount": round(float(row["refund_amount"]), 2)}
            for row in regional.to_dict(orient="records")
        ]

    payment_lag = None
    negative_payment_lags = None
    if "付款时间" in cleaned.columns:
        paired = cleaned[cleaned["付款时间"].notna() & cleaned["下单时间"].notna()].copy()
        lag_hours = (paired["付款时间"] - paired["下单时间"]).dt.total_seconds() / 3600
        negative_payment_lags = int((lag_hours < 0).sum())
        nonnegative = lag_hours[lag_hours >= 0]
        payment_lag = {
            "orders_with_both_timestamps": int(len(paired)),
            "negative_lags": negative_payment_lags,
            "median_hours_nonnegative": round(float(nonnegative.median()), 2) if not nonnegative.empty else None,
            "p90_hours_nonnegative": round(float(nonnegative.quantile(.9)), 2) if not nonnegative.empty else None,
        }

    daily_records = daily.to_dict(orient="records")
    kpi = {
        "orders": total_orders,
        "order_amount": order_amount,
        "orders_with_payment_timestamp": int(has_payment_time.sum()),
        "orders_with_positive_actual_payment": int(positive_paid_mask.sum()),
        "paid_amount": paid_amount,
        "payment_timestamp_rate": round(float(has_payment_time.sum()) / total_orders, 4) if total_orders else None,
        "positive_actual_payment_order_rate": round(float(positive_paid_mask.sum()) / total_orders, 4) if total_orders else None,
        "positive_actual_payment_aov": round(paid_amount / int(positive_paid_mask.sum()), 2) if positive_paid_mask.any() else None,
        "refund_orders": int(refund_mask.sum()) if refund_available else None,
        "refund_amount": refund_amount if refund_available else None,
        "refund_amount_to_order_amount": round(refund_amount / order_amount, 4) if refund_available and order_amount else None,
        "first_order_date": dated["下单时间"].min().strftime("%Y-%m-%d") if not dated.empty else None,
        "last_order_date": dated["下单时间"].max().strftime("%Y-%m-%d") if not dated.empty else None,
    }
    region_status = {"available": bool(region), "reason": None if region else "未提供有效收货地区字段。"}
    payment_status = {"available": "付款时间" in cleaned.columns and cleaned["付款时间"].notna().any(),
                      "reason": None if "付款时间" in cleaned.columns and cleaned["付款时间"].notna().any() else "未提供有效付款时间。"}
    limitations = [
        "公开镜像将该文件标注为天猫订单数据，但原始发布页和授权状态未能独立核实；结果只作项目测试，不代表当前经营情况。",
        "订单金额、实付金额和退款金额沿用来源字段；不将三者关系解释为利润、净收入或退款因果。",
        "支付订单定义为实付金额大于 0；这不是来源提供的独立订单状态字段。",
        "没有商品、客户、数量、渠道、流量或履约字段，因此商品分析、RFM、转化漏斗和履约指标不可用。",
        "下单日期集中于有限历史窗口，趋势只能描述样本期内波动，不能外推当前经营表现。",
    ]
    return {
        "dataset": {"type": "order_level_public_export", "provenance_status": "public mirror; source and license not independently verified"},
        "definitions": {
            "orders": "按订单号去重后的订单数。",
            "orders_with_payment_timestamp": "付款时间非空的去重订单数；以来源时间戳识别付款事件。",
            "payment_timestamp_rate": "付款时间非空的订单数 / 去重订单数；不代表订单完成率。",
            "orders_with_positive_actual_payment": "买家实付金额大于 0 的订单数；与付款时间口径分开呈现。",
            "positive_actual_payment_order_rate": "实付金额大于 0 的订单数 / 去重订单数；仅为字段口径比值。",
            "positive_actual_payment_aov": "实付金额合计 / 实付金额大于 0 的订单数。",
            "refund_orders": "退款金额大于 0 的订单数；仅在退款金额字段存在时计算。",
            "refund_amount_to_order_amount": "退款金额合计 / 订单金额合计；仅作字段比值，不代表退款率、损失率或利润率。",
            "payment_lag": "付款时间 - 下单时间，仅对两个时间均存在且非负的订单汇总。",
        },
        "kpi": kpi,
        "daily": daily_records,
        "region": region,
        "payment_lag": payment_lag,
        "availability": {
            "date_trend": {"available": not daily.empty, "reason": None if not daily.empty else "下单时间不可用。"},
            "region": region_status,
            "payment_timing": payment_status,
            "product": {"available": False, "reason": "未提供商品 ID、名称、品类或数量。"},
            "customer_rfm": {"available": False, "reason": "未提供用户 ID。"},
            "traffic_conversion": {"available": False, "reason": "未提供曝光、访客、加购等流量数据。"},
            "fulfilment": {"available": False, "reason": "未提供发货、签收和时效字段。"},
        },
        "validation": {
            "unique_order_ids": total_orders == len(cleaned),
            "daily_orders_reconcile": int(daily["orders"].sum()) == int(dated["订单号"].nunique()),
            "payment_timestamp_orders_reconcile": int(daily["payment_timestamp_orders"].sum()) == int(has_payment_time.loc[dated.index].sum()),
            "daily_order_amount_reconciles": abs(float(daily["order_amount"].sum()) - _money_sum(dated["订单金额"].fillna(0))) < 0.011,
            "daily_paid_amount_reconciles": abs(float(daily["paid_amount"].sum()) - _money_sum(dated["实付金额"].fillna(0))) < 0.011,
            "region_orders_reconcile_when_available": None if not region else sum(item["orders"] for item in region) == int(cleaned.loc[cleaned["收货省份"].notna(), "订单号"].nunique()),
        },
        "limitations": limitations,
    }


def _write_json(path: Path, value: Any) -> None:
    def convert(item: Any) -> Any:
        if item is pd.NA or item is pd.NaT:
            return None
        if hasattr(item, "item"):
            return item.item()
        if hasattr(item, "isoformat"):
            return item.isoformat()
        raise TypeError(f"Unsupported JSON value: {type(item).__name__}")

    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=convert), encoding="utf-8")


def _write_table(sheet, headers: list[str], rows: list[dict[str, Any]]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="175C4C")
        cell.font = Font(color="FFFFFF", bold=True)
    for row in rows:
        sheet.append([row.get(header) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cells in sheet.columns:
        width = min(max(max(len(str(cell.value or "")) for cell in cells) + 2, 12), 32)
        sheet.column_dimensions[cells[0].column_letter].width = width
        for cell in cells:
            cell.alignment = Alignment(vertical="top")


def write_order_level_outputs(output_dir: Path, cleaned: pd.DataFrame, cleaning_log: dict[str, Any], analysis: dict[str, Any]) -> None:
    """Write analysis files only beneath the caller-provided output directory."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    cleaned.to_csv(output_dir / "cleaned_orders.csv", index=False, encoding="utf-8-sig", date_format="%Y-%m-%d %H:%M:%S")
    _write_json(output_dir / "cleaning_log.json", cleaning_log)
    _write_json(output_dir / "analysis.json", analysis)
    workbook = Workbook()
    details = workbook.active
    details.title = "清洗订单"
    detail_rows = cleaned.astype(object).where(pd.notna(cleaned), None).to_dict(orient="records")
    _write_table(details, list(cleaned.columns), detail_rows)
    kpi = workbook.create_sheet("核心指标")
    _write_table(kpi, ["指标", "结果", "口径"], [
        {"指标": "订单数", "结果": analysis["kpi"]["orders"], "口径": analysis["definitions"]["orders"]},
        {"指标": "订单金额合计", "结果": analysis["kpi"]["order_amount"], "口径": "来源字段总金额汇总"},
        {"指标": "有付款时间订单数", "结果": analysis["kpi"]["orders_with_payment_timestamp"], "口径": analysis["definitions"]["orders_with_payment_timestamp"]},
        {"指标": "实付金额大于 0 的订单数", "结果": analysis["kpi"]["orders_with_positive_actual_payment"], "口径": analysis["definitions"]["orders_with_positive_actual_payment"]},
        {"指标": "实付金额合计", "结果": analysis["kpi"]["paid_amount"], "口径": "来源字段买家实际支付金额汇总"},
        {"指标": "付款时间订单占比", "结果": analysis["kpi"]["payment_timestamp_rate"], "口径": analysis["definitions"]["payment_timestamp_rate"]},
        {"指标": "实付金额大于 0 订单占比", "结果": analysis["kpi"]["positive_actual_payment_order_rate"], "口径": analysis["definitions"]["positive_actual_payment_order_rate"]},
        {"指标": "退款金额合计", "结果": analysis["kpi"]["refund_amount"], "口径": "来源字段退款金额汇总；不是利润或损失"},
    ])
    _write_table(workbook.create_sheet("日趋势"), ["下单日期", "orders", "order_amount", "payment_timestamp_orders", "positive_actual_payment_orders", "paid_amount", "refund_amount"], analysis["daily"])
    _write_table(workbook.create_sheet("地区汇总"), ["收货省份", "orders", "order_amount", "paid_amount", "refund_amount"], analysis["region"])
    limits = workbook.create_sheet("口径与限制")
    validation_passed = all(value is not False for value in analysis["validation"].values())
    availability_rows = [
        {"类型": f"不可用维度：{name}", "内容": item["reason"]}
        for name, item in analysis["availability"].items() if not item["available"]
    ]
    validation_rows = [
        {"类型": "校验结果", "内容": f"{name}: {'通过' if value is True else '未通过' if value is False else '不适用'}"}
        for name, value in analysis["validation"].items()
    ]
    _write_table(limits, ["类型", "内容"],
                 [{"类型": "口径", "内容": value} for value in analysis["definitions"].values()]
                 + [{"类型": "清洗决策", "内容": value} for value in cleaning_log["rules"]]
                 + validation_rows + availability_rows
                 + [{"类型": "限制", "内容": value} for value in analysis["limitations"]])
    workbook.save(output_dir / "order_analysis.xlsx")
    kpi_rows = "".join(f"<tr><th>{html.escape(str(key))}</th><td>{html.escape(str(value))}</td></tr>" for key, value in analysis["kpi"].items())
    daily_rows = "".join("<tr>" + "".join(f"<td>{html.escape(str(value))}</td>" for value in row.values()) + "</tr>" for row in analysis["daily"])
    region_rows = "".join("<tr>" + "".join(f"<td>{html.escape(str(value))}</td>" for value in row.values()) + "</tr>" for row in analysis["region"][:15])
    limitations = "".join(f"<li>{html.escape(value)}</li>" for value in analysis["limitations"])
    cleaning_items = "".join(f"<li>{html.escape(rule)}</li>" for rule in cleaning_log["rules"])
    unavailable_items = "".join(f"<li>{html.escape(name)}：{html.escape(item['reason'])}</li>"
                                 for name, item in analysis["availability"].items() if not item["available"])
    validation_items = "".join(f"<li>{html.escape(name)}：{'通过' if value is True else '未通过' if value is False else '不适用'}</li>"
                               for name, value in analysis["validation"].items())
    document = f"""<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>订单级数据分析</title>
<style>body{{font:15px/1.6 'Microsoft YaHei',sans-serif;max-width:1100px;margin:32px auto;padding:0 20px;color:#24332e}}h1,h2{{color:#175c4c}}table{{border-collapse:collapse;width:100%;margin:12px 0 28px}}th,td{{border:1px solid #d7ddd9;padding:7px;text-align:left}}th{{background:#edf4f0}}li{{margin:6px 0}}</style>
<h1>订单级数据分析报告</h1><p>数据为订单粒度公开镜像测试；未进行商品、用户、流量或履约推算。</p>
<h2>核心指标</h2><table>{kpi_rows}</table><h2>日趋势</h2><table><thead><tr>{''.join(f'<th>{html.escape(c)}</th>' for c in (analysis['daily'][0].keys() if analysis['daily'] else []))}</tr></thead><tbody>{daily_rows}</tbody></table>
<h2>地区分布（前 15）</h2><table><thead><tr>{''.join(f'<th>{html.escape(c)}</th>' for c in (analysis['region'][0].keys() if analysis['region'] else []))}</tr></thead><tbody>{region_rows}</tbody></table>
<h2>清洗记录</h2><p>原始 {cleaning_log['raw_rows']:,} 行，去重后 {cleaning_log['rows_after_deduplication']:,} 行，移除重复订单 {cleaning_log['duplicate_rows_removed']:,} 行。</p><ul>{cleaning_items}</ul>
<h2>不可用维度</h2><ul>{unavailable_items}</ul><h2>校验</h2><p>{'全部通过' if validation_passed else '存在未通过项'}</p><ul>{validation_items}</ul>
<h2>口径与局限</h2><ul>{limitations}</ul></html>"""
    (output_dir / "report.html").write_text(document, encoding="utf-8")
    validation = {"status": "passed" if all(value is not False for value in analysis["validation"].values()) else "failed",
                  "checks": analysis["validation"]}
    _write_json(output_dir / "validation.json", validation)
